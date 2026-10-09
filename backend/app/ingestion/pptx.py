from pathlib import Path
from app.ingestion.base import ParsedUnit


def parse_pptx(path: Path) -> list[ParsedUnit]:
    from pptx import Presentation
    units: list[ParsedUnit] = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if slide.has_notes_slide:
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        text = "\n".join(p.strip() for p in parts if p.strip())
        if text:
            units.append(ParsedUnit("text", text, slide=i))
        # TODO: extract pictures/diagrams -> Vision Agent
    return units
